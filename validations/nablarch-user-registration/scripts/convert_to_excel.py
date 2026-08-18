"""
把 5 阶段生成的成果物灌入 Excel 模版:
  - 02_Screen_Item_List.csv      -> [項目定義] Sheet
  - 03_Test_Cases.csv             -> [単体テスト仕様書] Sheet
  - 01_External_Design_W11AC01.md -> [画面レイアウト] Sheet (JSP 结构)

最终输出: W11AC01_ユーザー登録機能_詳細設計書.xlsx
"""
import csv
import os
import re
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ===== 路径 =====
BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(BASE, "templates", "Design_Doc_Template.xlsx")
CSV_ITEMS = os.path.join(BASE, "02_Screen_Item_List.csv")
CSV_TESTS = os.path.join(BASE, "03_Test_Cases.csv")
MD_DESIGN = os.path.join(BASE, "docs", "01_External_Design_W11AC01.md")
OUTPUT = os.path.join(BASE, "W11AC01_ユーザー登録機能_詳細設計書.xlsx")

# ===== 样式 =====
BODY_FONT = Font(name="Yu Gothic", size=10)
TITLE_FONT = Font(name="Yu Gothic", size=12, bold=True, color="1F3864")
CODE_FONT = Font(name="Consolas", size=9)
SUB_FILL = PatternFill("solid", fgColor="D9E1F2")
BORDER = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
TOP_LEFT = Alignment(horizontal="left", vertical="top", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_csv_to_sheet(ws, csv_path, start_row=4):
    """把 CSV 内容写入指定 Sheet,从 start_row 开始"""
    with open(csv_path, encoding="utf-8") as f:
        reader = csv.reader(f)
        rows = list(reader)
    if not rows:
        return 0

    # CSV 第一行是表头,跳过(Excel 已有表头)
    data_rows = rows[1:]
    for r_idx, row in enumerate(data_rows):
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=start_row + r_idx, column=c_idx, value=val)
            cell.font = BODY_FONT
            cell.alignment = LEFT if c_idx >= 3 else CENTER
            cell.border = BORDER
    return len(data_rows)


def extract_jsp_blocks(md_path):
    """从外部设计 MD 提取 jsp 代码块和标题"""
    with open(md_path, encoding="utf-8") as f:
        content = f.read()

    blocks = []
    # 匹配 ## 标题 + 后续 ```jsp 代码块
    pattern = re.compile(
        r"##\s+(.+?)\n.*?```jsp\n(.*?)```",
        re.DOTALL
    )
    for m in pattern.finditer(content):
        title = m.group(1).strip()
        code = m.group(2).strip()
        blocks.append((title, code))
    return blocks


def write_layout_to_sheet(ws, md_path, start_row=7):
    """把 MD 里的 JSP 代码块写入画面レイアウト Sheet"""
    blocks = extract_jsp_blocks(md_path)
    if not blocks:
        # 兜底:如果没有 jsp 代码块,至少写一段说明
        ws.cell(row=start_row, column=1,
                value="(JSP コードブロックが抽出できませんでした)").font = BODY_FONT
        return

    cur = start_row
    for title, code in blocks:
        # 子标题
        ws.merge_cells(start_row=cur, start_column=1, end_row=cur, end_column=8)
        cell = ws.cell(row=cur, column=1, value=f"■ {title}")
        cell.font = TITLE_FONT
        cell.fill = SUB_FILL
        cell.alignment = LEFT
        ws.row_dimensions[cur].height = 22
        cur += 1

        # 代码内容(按行写入,合并 A-H 列)
        lines = code.split("\n")
        for line in lines:
            ws.merge_cells(start_row=cur, start_column=1, end_row=cur, end_column=8)
            c = ws.cell(row=cur, column=1, value=line if line else " ")
            c.font = CODE_FONT
            c.alignment = TOP_LEFT
            c.border = BORDER
            cur += 1
        cur += 1  # 空行分隔


def main():
    print(f"加载模版: {TEMPLATE}")
    wb = load_workbook(TEMPLATE)
    print(f"Sheets: {wb.sheetnames}")

    # 1. 項目定義 Sheet <- 02_Screen_Item_List.csv
    print("\n[1/3] 灌入 項目定義 Sheet (02_Screen_Item_List.csv)")
    ws_items = wb["項目定義"]
    n = write_csv_to_sheet(ws_items, CSV_ITEMS, start_row=4)
    print(f"  ✓ 写入 {n} 行")
    # 调整行高
    for r in range(4, 4 + n):
        ws_items.row_dimensions[r].height = 36

    # 2. 単体テスト仕様書 Sheet <- 03_Test_Cases.csv
    print("\n[2/3] 灌入 単体テスト仕様書 Sheet (03_Test_Cases.csv)")
    ws_tests = wb["単体テスト仕様書"]
    n = write_csv_to_sheet(ws_tests, CSV_TESTS, start_row=4)
    print(f"  ✓ 写入 {n} 行")
    for r in range(4, 4 + n):
        ws_tests.row_dimensions[r].height = 42

    # 3. 画面レイアウト Sheet <- 01_External_Design_W11AC01.md (JSP 块)
    print("\n[3/3] 灌入 画面レイアウト Sheet (01_External_Design_W11AC01.md 的 JSP)")
    ws_layout = wb["画面レイアウト"]
    write_layout_to_sheet(ws_layout, MD_DESIGN, start_row=7)
    print("  ✓ JSP 代码块写入完毕")

    # 保存
    wb.save(OUTPUT)
    print(f"\n=== 完成 ===")
    print(f"输出: {OUTPUT}")
    print(f"大小: {os.path.getsize(OUTPUT)} bytes")

    # 验证
    print("\n=== 验证 ===")
    wb2 = load_workbook(OUTPUT)
    for name in wb2.sheetnames:
        ws = wb2[name]
        print(f"  [{name}] {ws.max_row} rows x {ws.max_column} cols")


if __name__ == "__main__":
    main()
