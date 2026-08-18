"""
生成 Nablarch 设计书 Excel 模版 (Design_Doc_Template.xlsx)
包含 3 个 Sheet:
  1. 画面レイアウト (画面布局预览)
  2. 項目定義 (项目/字段定义表)
  3. 単体テスト仕様書 (测试矩阵)
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os

# 样式
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(name="Yu Gothic", size=11, bold=True, color="FFFFFF")
SUB_FILL = PatternFill("solid", fgColor="D9E1F2")
SUB_FONT = Font(name="Yu Gothic", size=11, bold=True, color="000000")
BODY_FONT = Font(name="Yu Gothic", size=10)
BORDER = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def header_row(ws, headers, row=1):
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


wb = Workbook()

# Sheet 1: 画面レイアウト
ws1 = wb.active
ws1.title = "画面レイアウト"
ws1.merge_cells("A1:H1")
t = ws1.cell(row=1, column=1, value="画面レイアウト (Screen Layout Preview)")
t.font = Font(name="Yu Gothic", size=14, bold=True, color="FFFFFF")
t.fill = PatternFill("solid", fgColor="2F5496")
t.alignment = CENTER
ws1.row_dimensions[1].height = 28

for i, (k, v) in enumerate([("画面ID", "W11AC01"), ("画面名", "ユーザー登録"), ("版数", "1.0"), ("作成日", "")]):
    r = 2 + i
    ws1.cell(row=r, column=1, value=k).font = SUB_FONT
    ws1.cell(row=r, column=1).fill = SUB_FILL
    ws1.cell(row=r, column=1).alignment = LEFT
    ws1.cell(row=r, column=1).border = BORDER
    ws1.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    c = ws1.cell(row=r, column=2, value=v)
    c.font = BODY_FONT
    c.alignment = LEFT
    c.border = BORDER

ws1.merge_cells("A6:H6")
lt = ws1.cell(row=6, column=1, value="▼ 画面布局预览 (HTML/JSP 结构)")
lt.font = SUB_FONT
lt.fill = SUB_FILL
lt.alignment = LEFT
ws1.row_dimensions[6].height = 22

for r in range(7, 32):
    for c in range(1, 9):
        cell = ws1.cell(row=r, column=c, value="")
        cell.border = BORDER
        cell.alignment = LEFT

widths(ws1, [16, 18, 18, 18, 14, 14, 14, 14])

# Sheet 2: 項目定義
ws2 = wb.create_sheet("項目定義")
ws2.merge_cells("A1:G1")
t2 = ws2.cell(row=1, column=1, value="項目定義 (Screen Item Definition)")
t2.font = Font(name="Yu Gothic", size=14, bold=True, color="FFFFFF")
t2.fill = PatternFill("solid", fgColor="2F5496")
t2.alignment = CENTER
ws2.row_dimensions[1].height = 28

header_row(ws2, ["No", "画面ID", "項目名 (論理名)", "物理名", "Java型", "桁数", "精査ルール (Validation)"], row=3)
ws2.row_dimensions[3].height = 24

for r in range(4, 34):
    for c in range(1, 8):
        cell = ws2.cell(row=r, column=c, value="")
        cell.font = BODY_FONT
        cell.alignment = LEFT if c >= 3 else CENTER
        cell.border = BORDER

widths(ws2, [6, 12, 22, 22, 12, 10, 40])

# Sheet 3: 単体テスト仕様書
ws3 = wb.create_sheet("単体テスト仕様書")
ws3.merge_cells("A1:G1")
t3 = ws3.cell(row=1, column=1, value="単体テスト仕様書 (Unit Test Specification)")
t3.font = Font(name="Yu Gothic", size=14, bold=True, color="FFFFFF")
t3.fill = PatternFill("solid", fgColor="2F5496")
t3.alignment = CENTER
ws3.row_dimensions[1].height = 28

header_row(ws3, ["Test No", "テスト目的", "区分 (正常/異常)", "入力データ", "期待結果", "メッセージID", "備考"], row=3)
ws3.row_dimensions[3].height = 24

for r in range(4, 34):
    for c in range(1, 8):
        cell = ws3.cell(row=r, column=c, value="")
        cell.font = BODY_FONT
        cell.alignment = LEFT if c in (2, 4, 5, 7) else CENTER
        cell.border = BORDER

widths(ws3, [10, 28, 14, 30, 30, 22, 20])

out_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "templates", "Design_Doc_Template.xlsx")
wb.save(out_path)
print(f"OK -> {out_path}")
print(f"Sheets: {wb.sheetnames}")
