#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""创建横展开报告模板 Excel 文件"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

TEMPLATE_PATH = "/workspace/validations/office-cli/test-data/template-demo/横展开报告模板.xlsx"

# 通用样式
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION_FILL = PatternFill("solid", fgColor="DDEBF7")
TABLE_HEADER_FILL = PatternFill("solid", fgColor="BDD7EE")
WHITE_FONT = Font(color="FFFFFF", bold=True, size=16)
BLUE_FONT = Font(color="1F4E78", bold=True, size=12)
BOLD_FONT = Font(bold=True)
NORMAL_FONT = Font(size=11)
THIN_BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)
WRAP_CENTER = Alignment(wrap_text=True, vertical='center')
WRAP_LEFT = Alignment(wrap_text=True, vertical='center', horizontal='left')

def set_cell(ws, ref, value, **kwargs):
    cell = ws[ref]
    cell.value = value
    if kwargs.get('bold'):
        cell.font = Font(bold=True, size=kwargs.get('size', 11),
                         color=kwargs.get('font_color', None) or kwargs.get('font.color', None))
    else:
        cell.font = Font(size=kwargs.get('size', 11))
    if 'fill' in kwargs:
        cell.fill = PatternFill("solid", fgColor=kwargs['fill'])
    elif kwargs.get('bgcolor'):
        cell.fill = PatternFill("solid", fgColor=kwargs['bgcolor'])
    if kwargs.get('font_color'):
        if 'bold' in kwargs:
            cell.font = Font(bold=kwargs['bold'], size=kwargs.get('size', 11), color=kwargs['font_color'])
        else:
            cell.font = Font(size=kwargs.get('size', 11), color=kwargs['font_color'])
    cell.alignment = kwargs.get('align', WRAP_LEFT)
    if kwargs.get('border'):
        cell.border = THIN_BORDER
    return cell

wb = Workbook()

# =====================================================
# Sheet 1: 横展开报告
# =====================================================
ws1 = wb.active
ws1.title = "1_横展开报告"
ws1.sheet_properties.tabColor = "1F4E78"

# 列宽
ws1.column_dimensions['A'].width = 18
ws1.column_dimensions['B'].width = 30
ws1.column_dimensions['C'].width = 35
ws1.column_dimensions['D'].width = 18
ws1.column_dimensions['E'].width = 28
ws1.column_dimensions['F'].width = 10

# 标题行
ws1.merge_cells('A1:F1')
c = set_cell(ws1, 'A1', "横展开报告", bold=True, size=16, font_color="FFFFFF", fill="1F4E78",
             align=Alignment(horizontal='center', vertical='center'))
ws1.row_dimensions[1].height = 40

# 报告信息
set_cell(ws1, 'A3', "报告编号：", bold=True)
set_cell(ws1, 'B3', "{{report_id}}")
set_cell(ws1, 'D3', "作成日：", bold=True)
set_cell(ws1, 'E3', "{{create_date}}")

# 1. 横展开目的
ws1.merge_cells('A5:F5')
set_cell(ws1, 'A5', "1. 横展开目的", bold=True, size=12, fill="DDEBF7", font_color="1F4E78")
ws1.merge_cells('A6:F7')
c = set_cell(ws1, 'A6', "{{purpose}}")
c.alignment = Alignment(wrap_text=True, vertical='top')
ws1.row_dimensions[6].height = 40

# 2. 横展开分支
ws1.merge_cells('A8:F8')
set_cell(ws1, 'A8', "2. 横展开分支（対象ブランチ）", bold=True, size=12, fill="DDEBF7", font_color="1F4E78")
set_cell(ws1, 'A9', "対象ブランチ：", bold=True)
set_cell(ws1, 'B9', "{{branches}}")
set_cell(ws1, 'A10', "コミットハッシュ：", bold=True)
set_cell(ws1, 'B10', "{{commit_hash}}")

# 3. 横展开作业时间
ws1.merge_cells('A12:F12')
set_cell(ws1, 'A12', "3. 横展开作业时间", bold=True, size=12, fill="DDEBF7", font_color="1F4E78")
set_cell(ws1, 'A13', "开始时间：", bold=True)
set_cell(ws1, 'B13', "{{start_time}}")
set_cell(ws1, 'D13', "结束时间：", bold=True)
set_cell(ws1, 'E13', "{{end_time}}")
set_cell(ws1, 'A14', "总工时：", bold=True)
set_cell(ws1, 'B14', "{{total_hours}} 小时")

# 4. 横展开方法
ws1.merge_cells('A16:F16')
set_cell(ws1, 'A16', "4. 横展开方法", bold=True, size=12, fill="DDEBF7", font_color="1F4E78")
ws1.merge_cells('A17:F18')
c = set_cell(ws1, 'A17', "{{method}}")
c.alignment = Alignment(wrap_text=True, vertical='top')
ws1.row_dimensions[17].height = 40

# 5. 横展开观点
ws1.merge_cells('A19:F19')
set_cell(ws1, 'A19', "5. 横展开观点（检查项目）", bold=True, size=12, fill="DDEBF7", font_color="1F4E78")

headers5 = [("A20", "No.", 6), ("B20", "检查观点", 28), ("C20", "检查说明", 65), ("D20", "优先级", 12)]
for ref, val, width in headers5:
    set_cell(ws1, ref, val, bold=True, fill="BDD7EE", border=True)

for i, vp in enumerate(["1", "2", "3"], 21):
    set_cell(ws1, f'A{i}', vp, border=True)
    set_cell(ws1, f'B{i}', "{{viewpoint_" + vp + "}}", border=True)
    set_cell(ws1, f'C{i}', "{{viewpoint_" + vp + "_desc}}", border=True)
    set_cell(ws1, f'D{i}', "{{viewpoint_" + vp + "_priority}}", border=True)

# 6. 横展开结论
ws1.merge_cells('A25:F25')
set_cell(ws1, 'A25', "6. 横展开结论", bold=True, size=12, fill="DDEBF7", font_color="1F4E78")
set_cell(ws1, 'A26', "总文件数：", bold=True)
set_cell(ws1, 'B26', "{{total_files}}")
set_cell(ws1, 'A27', "检查总项目数：", bold=True)
set_cell(ws1, 'B27', "{{total_checks}}")
set_cell(ws1, 'A28', "问题数：", bold=True)
set_cell(ws1, 'B28', "{{issue_count}}")
set_cell(ws1, 'D28', "横展开状态：", bold=True)
set_cell(ws1, 'E28', "{{status}}")
set_cell(ws1, 'A29', "结论总结：", bold=True)
ws1.merge_cells('B29:F30')
c = set_cell(ws1, 'B29', "{{conclusion}}")
c.alignment = Alignment(wrap_text=True, vertical='top')
ws1.row_dimensions[29].height = 50

# 签名
set_cell(ws1, 'A32', "作成者：", bold=True)
set_cell(ws1, 'B32', "{{author}}")
set_cell(ws1, 'D32', "承认者：", bold=True)
set_cell(ws1, 'E32', "{{approver}}")


# =====================================================
# Sheet 2: 监测一览
# =====================================================
ws2 = wb.create_sheet("2_监测一览")
ws2.sheet_properties.tabColor = "4472C4"

ws2.column_dimensions['A'].width = 6
ws2.column_dimensions['B'].width = 30
ws2.column_dimensions['C'].width = 14
ws2.column_dimensions['D'].width = 14
ws2.column_dimensions['E'].width = 14
ws2.column_dimensions['F'].width = 14
ws2.column_dimensions['G'].width = 14
ws2.column_dimensions['H'].width = 30

# 标题
ws2.merge_cells('A1:H1')
c = set_cell(ws2, 'A1', "コード検索・モニタリング結果一覧（监测一览）", bold=True, size=14,
             font_color="FFFFFF", fill="4472C4", align=Alignment(horizontal='center', vertical='center'))
ws2.row_dimensions[1].height = 35

# 汇总信息
set_cell(ws2, 'A3', "报告编号：", bold=True)
set_cell(ws2, 'B3', "{{report_id}}")
set_cell(ws2, 'D3', "対象ブランチ：", bold=True)
set_cell(ws2, 'E3', "{{branches}}")
set_cell(ws2, 'G3', "集计日：", bold=True)
set_cell(ws2, 'H3', "{{create_date}}")

# 表头
headers2 = [
    ("A5", "No.", 6), ("B5", "ファイル名 (文件)", 30),
    ("C5", "观点1 該当数", 12), ("D5", "观点2 該当数", 12),
    ("E5", "观点3 該当数", 12), ("F5", "合計", 10),
    ("G5", "ステータス", 12), ("H5", "備考", 30)
]
for ref, val, width in headers2:
    set_cell(ws2, ref, val, bold=True, fill="D6E4F0", font_color="1F4E78", border=True,
             align=Alignment(horizontal='center', vertical='center', wrap_text=True))

for i in range(1, 11):
    row = 5 + i
    set_cell(ws2, f'A{row}', f"{{% file_{i}_no %}}" if False else f"{{file_{i}_no}}", border=True)
    set_cell(ws2, f'B{row}', f"{{file_{i}_name}}", border=True)
    set_cell(ws2, f'C{row}', f"{{file_{i}_vp1}}", border=True)
    set_cell(ws2, f'D{row}', f"{{file_{i}_vp2}}", border=True)
    set_cell(ws2, f'E{row}', f"{{file_{i}_vp3}}", border=True)
    set_cell(ws2, f'F{row}', f"{{file_{i}_total}}", border=True)
    set_cell(ws2, f'G{row}', f"{{file_{i}_status}}", border=True)
    set_cell(ws2, f'H{row}', f"{{file_{i}_remark}}", border=True)
ws2.row_dimensions[5].height = 30

# 汇总行
sum_row = 16
set_cell(ws2, f'A{sum_row}', "", border=True)
ws2.merge_cells(f'A{sum_row}:B{sum_row}')
c = set_cell(ws2, f'A{sum_row}', "合計", bold=True, fill="FFF2CC", border=True,
             align=Alignment(horizontal='center', vertical='center'))
for col in ['C', 'D', 'E', 'F']:
    set_cell(ws2, f'{col}{sum_row}', f"{{total_{col.lower()}}}", bold=True, fill="FFF2CC", border=True)
set_cell(ws2, f'G{sum_row}', "", fill="FFF2CC", border=True)
set_cell(ws2, f'H{sum_row}', "", fill="FFF2CC", border=True)


# =====================================================
# Sheet 3: 详细明細
# =====================================================
ws3 = wb.create_sheet("3_详细明細")
ws3.sheet_properties.tabColor = "ED7D31"

ws3.column_dimensions['A'].width = 6
ws3.column_dimensions['B'].width = 32
ws3.column_dimensions['C'].width = 10
ws3.column_dimensions['D'].width = 10
ws3.column_dimensions['E'].width = 55
ws3.column_dimensions['F'].width = 18
ws3.column_dimensions['G'].width = 14
ws3.column_dimensions['H'].width = 10
ws3.column_dimensions['I'].width = 30

# 标题
ws3.merge_cells('A1:I1')
c = set_cell(ws3, 'A1', "コード検索詳細明細（详细检测结果）", bold=True, size=14,
             font_color="FFFFFF", fill="ED7D31", align=Alignment(horizontal='center', vertical='center'))
ws3.row_dimensions[1].height = 35

# 汇总信息
set_cell(ws3, 'A3', "报告编号：", bold=True)
set_cell(ws3, 'B3', "{{report_id}}")
set_cell(ws3, 'D3', "检查观点：", bold=True)
ws3.merge_cells('E3:H3')
set_cell(ws3, 'E3', "{{viewpoints_summary}}")

# 表头
headers3 = [
    ("A5", "No.", 6), ("B5", "ファイルパス (文件路径)", 36),
    ("C5", "行番号", 9), ("D5", "列番号", 9),
    ("E5", "コード内容 (代码片段)", 60), ("F5", "该当検査観点", 18),
    ("G5", "重要度", 12), ("H5", "要修正", 9),
    ("I5", "修正内容 / コメント", 30)
]
for ref, val, width in headers3:
    set_cell(ws3, ref, val, bold=True, fill="FBE5D6", font_color="843C0C", border=True,
             align=Alignment(horizontal='center', vertical='center', wrap_text=True))
ws3.row_dimensions[5].height = 35

# 占位示例数据（20 行）
for i in range(1, 21):
    row = 5 + i
    set_cell(ws3, f'A{row}', f"{{detail_{i}_no}}", border=True)
    set_cell(ws3, f'B{row}', f"{{detail_{i}_file}}", border=True)
    set_cell(ws3, f'C{row}', f"{{detail_{i}_line}}", border=True)
    set_cell(ws3, f'D{row}', f"{{detail_{i}_col}}", border=True)
    set_cell(ws3, f'E{row}', f"{{detail_{i}_code}}", border=True)
    set_cell(ws3, f'F{row}', f"{{detail_{i}_viewpoint}}", border=True)
    set_cell(ws3, f'G{row}', f"{{detail_{i}_severity}}", border=True)
    set_cell(ws3, f'H{row}', f"{{detail_{i}_needfix}}", border=True)
    set_cell(ws3, f'I{row}', f"{{detail_{i}_comment}}", border=True)

wb.save(TEMPLATE_PATH)
print(f"✅ 模板已生成: {TEMPLATE_PATH}")
print(f"   - Sheet1: {ws1.title} (横展开主报告)")
print(f"   - Sheet2: {ws2.title} (监测一览)")
print(f"   - Sheet3: {ws3.title} (详细明細)")
